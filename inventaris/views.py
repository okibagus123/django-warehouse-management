from decimal import Decimal, InvalidOperation
from django.contrib import messages
from django.contrib.auth import (
    authenticate,
    login,
    logout
)
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.core.paginator import Paginator  # <-- Ditambahkan untuk Pagination
from django.db import (
    IntegrityError,
    transaction
)
from django.db.models import (
    F,
    Sum,
    Count,
    Q  # <-- Ditambahkan untuk filter pencarian (Search)
)
from django.shortcuts import (
    render,
    redirect,
    get_object_or_404
)
from openpyxl import load_workbook
from .models import (
    Kategori,
    Supplier,
    Lokasi,
    Barang,
    Inventory,
    StockMovement,
    UserProfile
)
from .forms import (
    BarangForm,
    KategoriForm,
    SupplierForm,
    LokasiForm,
    BarangMasukForm,
    BarangKeluarForm,
    TransferForm,
    UserCreateForm,
    UserEditForm,
    UserProfileForm
)
from django.utils import timezone
from datetime import timedelta
from django.db.models.functions import TruncDate
import json
import csv
from django.http import HttpResponse


# =========================
# DASHBOARD
# =========================
@login_required(login_url='/login/')
def dasbor_view(request):

    
    context = {

        'jumlah_sku': Barang.objects.count(),

        'total_masuk': StockMovement.objects.filter(
            tipe="MASUK"
        ).count(),

        'total_keluar': StockMovement.objects.filter(
            tipe="KELUAR"
        ).count(),

        'jumlah_lokasi': Lokasi.objects.count(),

    }


    return render(
            request,
            'inventaris/dashboard.html'
    )



# =========================
# MASTER MENU
# =========================
@login_required(login_url='/login/')
def master_data_view(request):

    return render(
        request,
        'inventaris/master/master.html'
    )


# ============================================================
# MASTER BARANG
# ============================================================

@login_required(login_url='/login/')
def master_barang_view(request):

    daftar_barang = Barang.objects.all().order_by('nama')

    return render(
        request,
        'inventaris/master/barang/master_barang.html',
        {
            'daftar_barang': daftar_barang
        }
    )


def tambah_barang_view(request):

    if request.method == "POST":

        form = BarangForm(request.POST)

        if form.is_valid():
            form.save()
            return redirect('master_barang')

    else:
        form = BarangForm()

    return render(
        request,
        'inventaris/master/barang/tambah_barang.html',
        {
            'form': form
        }
    )

def edit_barang_view(request, id):

    barang = get_object_or_404(
        Barang,
        id=id
    )


    if request.method == "POST":

        form = BarangForm(
            request.POST,
            instance=barang
        )


        if form.is_valid():

            form.save()


            messages.success(
                request,
                "Barang berhasil diperbarui"
            )


            return redirect('master_barang')


        else:

            messages.error(
                request,
                "Gagal memperbarui barang. Silakan periksa kembali data yang dimasukkan."
            )


    else:

        form = BarangForm(
            instance=barang
        )


    return render(
        request,
        'inventaris/master/barang/edit_barang.html',
        {
            'form': form,
            'barang': barang
        }
    )

def hapus_barang_view(request, id):

    barang = get_object_or_404(
        Barang,
        id=id
    )

    if request.method == "POST":
        barang.delete()

    return redirect('master_barang')


def hapus_semua_barang_view(request):

    if request.method == "POST":
        Barang.objects.all().delete()

    return redirect('master_barang')



# ============================================================
# MASTER KATEGORI
# ============================================================
@login_required(login_url='/login/')
def master_kategori_view(request):

    data_kategori = Kategori.objects.all().order_by('nama')

    return render(
        request,
        'inventaris/master/kategori/master_kategori.html',
        {
            'data_kategori': data_kategori
        }
    )


def tambah_kategori_view(request):

    if request.method == "POST":

        form = KategoriForm(request.POST)

        if form.is_valid():
            form.save()
            return redirect('master_kategori')

    else:
        form = KategoriForm()

    return render(
        request,
        'inventaris/master/kategori/tambah_kategori.html',
        {
            'form': form
        }
    )


def edit_kategori_view(request, id):

    kategori = get_object_or_404(
        Kategori,
        id=id
    )

    if request.method == "POST":

        form = KategoriForm(
            request.POST,
            instance=kategori
        )

        if form.is_valid():
            form.save()
            return redirect('master_kategori')

    else:

        form = KategoriForm(
            instance=kategori
        )

    return render(
        request,
        'inventaris/master/kategori/edit_kategori.html',
        {
            'form': form,
            'kategori': kategori
        }
    )


def hapus_kategori_view(request, id):

    kategori = get_object_or_404(
        Kategori,
        id=id
    )

    if request.method == "POST":
        kategori.delete()

    return redirect('master_kategori')


def hapus_semua_kategori_view(request):

    if request.method == "POST":
        Kategori.objects.all().delete()

    return redirect('master_kategori')



# ============================================================
# MASTER SUPPLIER
# ============================================================
@login_required(login_url='/login/')
def master_supplier_view(request):

    data_supplier = Supplier.objects.all().order_by('nama')


    return render(
        request,
        'inventaris/master/supplier/master_supplier.html',
        {
            'data_supplier': data_supplier
        }
    )





def tambah_supplier_view(request):


    if request.method == "POST":


        form = SupplierForm(request.POST)



        if form.is_valid():


            form.save()


            messages.success(
                request,
                "Supplier berhasil ditambahkan"
            )


            return redirect('master_supplier')



        else:


            messages.error(
                request,
                "Supplier gagal ditambahkan. Periksa kembali data."
            )



    else:


        form = SupplierForm()



    return render(
        request,
        'inventaris/master/supplier/tambah_supplier.html',
        {
            'form': form
        }
    )







def edit_supplier_view(request, id):


    supplier = get_object_or_404(
        Supplier,
        id=id
    )



    if request.method == "POST":



        form = SupplierForm(
            request.POST,
            instance=supplier
        )



        if form.is_valid():



            form.save()



            messages.success(
                request,
                "Supplier berhasil diperbarui"
            )



            return redirect('master_supplier')




        else:



            messages.error(
                request,
                "Supplier gagal diperbarui. Periksa kembali data."
            )





    else:



        form = SupplierForm(
            instance=supplier
        )





    return render(
        request,
        'inventaris/master/supplier/edit_supplier.html',
        {
            'form': form,
            'supplier': supplier
        }
    )








def hapus_supplier_view(request, id):


    supplier = get_object_or_404(
        Supplier,
        id=id
    )



    if request.method == "POST":


        nama_supplier = supplier.nama


        supplier.delete()



        messages.success(
            request,
            f"Supplier {nama_supplier} berhasil dihapus"
        )



    return redirect('master_supplier')










def hapus_semua_supplier_view(request):


    if request.method == "POST":


        jumlah = Supplier.objects.count()



        Supplier.objects.all().delete()



        messages.success(
            request,
            f"{jumlah} data supplier berhasil dihapus"
        )



    return redirect('master_supplier')



# ============================================================
# MASTER LOKASI
# ============================================================

def master_lokasi_view(request):

    data_lokasi = Lokasi.objects.all().order_by('kode_lokasi')

    return render(
        request,
        'inventaris/master/lokasi/master_lokasi.html',
        {
            'data_lokasi': data_lokasi
        }
    )


def tambah_lokasi_view(request):

    if request.method == "POST":

        form = LokasiForm(request.POST)

        if form.is_valid():
            form.save()
            return redirect('master_lokasi')

    else:
        form = LokasiForm()

    return render(
        request,
        'inventaris/master/lokasi/tambah_lokasi.html',
        {
            'form': form
        }
    )


def edit_lokasi_view(request, id):

    lokasi = get_object_or_404(
        Lokasi,
        id=id
    )

    if request.method == "POST":

        form = LokasiForm(
            request.POST,
            instance=lokasi
        )

        if form.is_valid():
            form.save()
            return redirect('master_lokasi')

    else:

        form = LokasiForm(
            instance=lokasi
        )

    return render(
        request,
        'inventaris/master/lokasi/edit_lokasi.html',
        {
            'form': form,
            'lokasi': lokasi
        }
    )


def hapus_lokasi_view(request, id):

    lokasi = get_object_or_404(
        Lokasi,
        id=id
    )

    if request.method == "POST":
        lokasi.delete()

    return redirect('master_lokasi')


def hapus_semua_lokasi_view(request):

    if request.method == "POST":
        Lokasi.objects.all().delete()

    return redirect('master_lokasi')

# ============================================================
# IMPORT EXCEL MASTER DATA
# ============================================================
@login_required(login_url='/login/')
def _excel_text(value):

    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()



def import_excel_view(request, jenis):

    konfigurasi = {

        'barang': {
            'judul': 'Barang',
            'redirect': 'master_barang',
            'kolom': (
                'sku, barcode, nama, kategori, supplier, '
                'satuan, minimum_stock, harga_satuan'
            )
        },

        'kategori': {
            'judul': 'Kategori',
            'redirect': 'master_kategori',
            'kolom': 'nama, deskripsi'
        },

        'supplier': {
            'judul': 'Supplier',
            'redirect': 'master_supplier',
            'kolom': 'nama, kontak, alamat'
        },

        'lokasi': {
            'judul': 'Lokasi',
            'redirect': 'master_lokasi',
            'kolom': 'kode_lokasi, nama_lokasi'
        },
    }


    if jenis not in konfigurasi:
        return redirect('master')


    config = konfigurasi[jenis]


    if request.method == "POST":

        file_excel = request.FILES.get('file')

        if not file_excel:

            messages.error(
                request,
                'Silakan pilih file Excel terlebih dahulu.'
            )

            return redirect(
                'import_excel',
                jenis=jenis
            )


        if not file_excel.name.lower().endswith('.xlsx'):

            messages.error(
                request,
                'File harus berformat .xlsx'
            )

            return redirect(
                'import_excel',
                jenis=jenis
            )


        try:

            workbook = load_workbook(
                file_excel,
                data_only=True
            )

            worksheet = workbook.active


            header = [
                _excel_text(cell.value)
                .lower()
                .replace(' ', '_')
                for cell in worksheet[1]
            ]


            berhasil = 0
            gagal = 0


            for values in worksheet.iter_rows(
                min_row=2,
                values_only=True
            ):

                if not any(
                    value is not None
                    for value in values
                ):
                    continue


                data = dict(
                    zip(
                        header,
                        values
                    )
                )


                try:

                    # ==========================================
                    # KATEGORI
                    # ==========================================

                    if jenis == 'kategori':

                        nama = _excel_text(
                            data.get('nama')
                        )

                        if not nama:
                            gagal += 1
                            continue


                        Kategori.objects.update_or_create(
                            nama=nama,
                            defaults={
                                'deskripsi':
                                    _excel_text(
                                        data.get('deskripsi')
                                    )
                                    or None
                            }
                        )


                    # ==========================================
                    # SUPPLIER
                    # ==========================================

                    elif jenis == 'supplier':

                        nama = _excel_text(
                            data.get('nama')
                        )

                        if not nama:
                            gagal += 1
                            continue


                        Supplier.objects.update_or_create(
                            nama=nama,
                            defaults={
                                'kontak':
                                    _excel_text(
                                        data.get('kontak')
                                    ),

                                'alamat':
                                    _excel_text(
                                        data.get('alamat')
                                    )
                                    or None
                            }
                        )


                    # ==========================================
                    # LOKASI
                    # ==========================================

                    elif jenis == 'lokasi':

                        kode = _excel_text(
                            data.get('kode_lokasi')
                        )

                        nama = _excel_text(
                            data.get('nama_lokasi')
                        )


                        if not kode or not nama:
                            gagal += 1
                            continue


                        Lokasi.objects.update_or_create(
                            kode_lokasi=kode,
                            defaults={
                                'nama_lokasi': nama
                            }
                        )


                    # ==========================================
                    # BARANG
                    # ==========================================

                    elif jenis == 'barang':

                        sku = _excel_text(
                            data.get('sku')
                        )

                        nama = _excel_text(
                            data.get('nama')
                        )


                        if not sku or not nama:
                            gagal += 1
                            continue


                        nama_kategori = _excel_text(
                            data.get('kategori')
                        )

                        nama_supplier = _excel_text(
                            data.get('supplier')
                        )


                        kategori = None

                        if nama_kategori:

                            kategori = (
                                Kategori.objects
                                .filter(
                                    nama__iexact=nama_kategori
                                )
                                .first()
                            )

                            if kategori is None:

                                kategori = Kategori.objects.create(
                                    nama=nama_kategori
                                )


                        supplier = None

                        if nama_supplier:

                            supplier = (
                                Supplier.objects
                                .filter(
                                    nama__iexact=nama_supplier
                                )
                                .first()
                            )

                            if supplier is None:

                                supplier = Supplier.objects.create(
                                    nama=nama_supplier
                                )


                        barcode = _excel_text(
                            data.get('barcode')
                        ) or None


                        satuan = _excel_text(
                            data.get('satuan')
                        ) or 'PCS'


                        try:

                            minimum_stock = int(
                                data.get('minimum_stock') or 5
                            )

                        except (ValueError, TypeError):

                            minimum_stock = 5


                        try:

                            harga_satuan = Decimal(
                                str(
                                    data.get('harga_satuan') or 0
                                )
                            )

                        except (
                            InvalidOperation,
                            ValueError,
                            TypeError
                        ):

                            harga_satuan = Decimal('0')


                        Barang.objects.update_or_create(

                            sku=sku,

                            defaults={

                                'barcode': barcode,

                                'nama': nama,

                                'kategori': kategori,

                                'supplier': supplier,

                                'satuan': satuan,

                                'minimum_stock':
                                    minimum_stock,

                                'harga_satuan':
                                    harga_satuan,
                            }
                        )


                    berhasil += 1


                except (
                    IntegrityError,
                    ValueError,
                    TypeError
                ):

                    gagal += 1


            messages.success(
                request,
                f'Import selesai. '
                f'{berhasil} data berhasil, '
                f'{gagal} data gagal.'
            )


            return redirect(
                config['redirect']
            )


        except Exception as error:

            messages.error(
                request,
                f'Gagal membaca Excel: {error}'
            )


    return render(
        request,
        'inventaris/master/import_excel.html',
        {
            'jenis': jenis,
            'judul': config['judul'],
            'kolom': config['kolom'],
        }
    )
# ========================================================
# DETAIL BARANG & LOKASI
# ========================================================

@login_required(login_url='/login/')
def detail_barang_view(request, pk):
    barang = get_object_or_404(Barang, pk=pk)
    
    # Ambil distribusi stok di berbagai lokasi
    stok_lokasi = Inventory.objects.filter(barang=barang).select_related('lokasi')
    total_stok = sum(item.jumlah for item in stok_lokasi)
    
    # Ambil riwayat pergerakan khusus barang ini
    riwayat = StockMovement.objects.filter(barang=barang).select_related(
        'user', 'lokasi_asal', 'lokasi_tujuan'
    ).order_by('-tanggal')[:50] # Batasi 50 aktivitas terakhir agar ringan

    return render(request, 'inventaris/master/detail_barang.html', {
        'barang': barang,
        'stok_lokasi': stok_lokasi,
        'total_stok': total_stok,
        'riwayat': riwayat
    })

@login_required(login_url='/login/')
def detail_lokasi_view(request, pk):
    lokasi = get_object_or_404(Lokasi, pk=pk)
    
    # Ambil daftar barang yang ada di lokasi ini (jumlah > 0)
    isi_lokasi = Inventory.objects.filter(lokasi=lokasi, jumlah__gt=0).select_related(
        'barang', 'barang__kategori'
    ).order_by('barang__nama')
    
    total_unit = sum(item.jumlah for item in isi_lokasi)

    return render(request, 'inventaris/master/detail_lokasi.html', {
        'lokasi': lokasi,
        'isi_lokasi': isi_lokasi,
        'total_unit': total_unit
    })

# =========================
# TRANSAKSI
# =========================

def transaksi_view(request):

    return render(
        request,
        'inventaris/transaksi/transaksi.html'
    )


# =========================
# TRANSAKSI BARANG MASUK
# =========================
@login_required(login_url='/login/')
def barang_masuk_view(request):


    if request.method == "POST":


        form = BarangMasukForm(
            request.POST
        )



        if form.is_valid():


            barang = form.cleaned_data['barang']

            lokasi = form.cleaned_data['lokasi']

            jumlah = form.cleaned_data['jumlah']

            keterangan = form.cleaned_data['keterangan']



            with transaction.atomic():


                inventory, created = Inventory.objects.get_or_create(

                    barang=barang,

                    lokasi=lokasi,

                    defaults={

                        'jumlah':0

                    }

                )



                inventory.jumlah += jumlah

                inventory.save()



                StockMovement.objects.create(

                    barang=barang,

                    tipe="MASUK",

                    jumlah=jumlah,

                    lokasi_tujuan=lokasi,

                    user=request.user,

                    keterangan=keterangan

                )



            messages.success(

                request,

                "Barang masuk berhasil diproses"

            )



            return redirect(

                'barang_masuk'

            )



    else:


        form = BarangMasukForm()



    return render(

        request,

        'inventaris/transaksi/barang_masuk.html',

        {

            'form':form

        }

    )


# =========================
# TRANSAKSI BARANG KELUAR
# =========================
@login_required(login_url='/login/')
def barang_keluar_view(request):


    if request.method == "POST":


        form = BarangKeluarForm(

            request.POST

        )



        if form.is_valid():


            barang = form.cleaned_data['barang']

            lokasi = form.cleaned_data['lokasi']

            jumlah = form.cleaned_data['jumlah']

            keterangan = form.cleaned_data['keterangan']



            with transaction.atomic():


                inventory = Inventory.objects.filter(

                    barang=barang,

                    lokasi=lokasi

                ).first()



                # Validasi stok tersedia

                if not inventory:


                    messages.error(

                        request,

                        "Stok barang tidak ditemukan pada lokasi tersebut"

                    )


                    return redirect(

                        'barang_keluar'

                    )



                # Validasi jumlah stok

                if inventory.jumlah < jumlah:


                    messages.error(

                        request,

                        "Jumlah barang keluar melebihi stok tersedia"

                    )


                    return redirect(

                        'barang_keluar'

                    )



                # Kurangi stok

                inventory.jumlah -= jumlah

                inventory.save()



                # Simpan histori transaksi

                StockMovement.objects.create(

                    barang=barang,

                    tipe="KELUAR",

                    jumlah=jumlah,

                    lokasi_asal=lokasi,

                    user=request.user,

                    keterangan=keterangan

                )



            messages.success(

                request,

                "Barang keluar berhasil diproses"

            )


            return redirect(

                'barang_keluar'

            )



    else:


        form = BarangKeluarForm()



    return render(

        request,

        'inventaris/transaksi/barang_keluar.html',

        {

            'form':form

        }

    )

# =========================
# TRANSAKSI TRANSFER
# =========================
@login_required(login_url='/login/')
def transfer_view(request):


    if request.method == "POST":


        form = TransferForm(
            request.POST
        )


        if form.is_valid():


            barang = form.cleaned_data['barang']

            lokasi_asal = form.cleaned_data['lokasi_asal']

            lokasi_tujuan = form.cleaned_data['lokasi_tujuan']

            jumlah = form.cleaned_data['jumlah']

            keterangan = form.cleaned_data['keterangan']



            if lokasi_asal == lokasi_tujuan:


                messages.error(

                    request,

                    "Lokasi asal dan tujuan tidak boleh sama"

                )


                return redirect(
                    'transfer'
                )



            with transaction.atomic():


                stok_asal = Inventory.objects.filter(

                    barang=barang,

                    lokasi=lokasi_asal

                ).first()



                if not stok_asal or stok_asal.jumlah < jumlah:


                    messages.error(

                        request,

                        "Stok lokasi asal tidak mencukupi"

                    )


                    return redirect(

                        'transfer'

                    )



                # Kurangi stok lokasi asal

                stok_asal.jumlah -= jumlah

                stok_asal.save()



                # Tambah stok lokasi tujuan

                stok_tujuan, created = Inventory.objects.get_or_create(

                    barang=barang,

                    lokasi=lokasi_tujuan,

                    defaults={

                        'jumlah':0

                    }

                )



                stok_tujuan.jumlah += jumlah

                stok_tujuan.save()



                # Simpan histori transfer

                StockMovement.objects.create(

                    barang=barang,

                    tipe="TRANSFER",

                    jumlah=jumlah,

                    lokasi_asal=lokasi_asal,

                    lokasi_tujuan=lokasi_tujuan,

                    user=request.user,

                    keterangan=keterangan

                )



            messages.success(

                request,

                "Transfer barang berhasil"

            )


            return redirect(

                'transfer'

            )



    else:


        form = TransferForm()



    return render(

        request,

        'inventaris/transaksi/transfer_barang.html',

        {

            'form':form

        }

    )

# ========================================================
# STOCK ADJUSTMENT (PENYESUAIAN STOK)
# ========================================================
@login_required(login_url='/login/')
def adjustment_stok_view(request):
    if request.method == 'POST':
        barang_id = request.POST.get('barang')
        lokasi_id = request.POST.get('lokasi')
        jumlah_fisik = int(request.POST.get('jumlah_fisik', 0))
        keterangan = request.POST.get('keterangan')

        barang = get_object_or_404(Barang, id=barang_id)
        lokasi = get_object_or_404(Lokasi, id=lokasi_id)

        try:
            with transaction.atomic():
                # Ambil data inventory saat ini (jika belum ada, buat dengan jumlah 0)
                inventory, created = Inventory.objects.get_or_create(
                    barang=barang,
                    lokasi=lokasi,
                    defaults={'jumlah': 0}
                )

                stok_sistem = inventory.jumlah
                selisih = jumlah_fisik - stok_sistem

                if selisih != 0:
                    # 1. Update stok inventory menjadi sesuai fisik
                    inventory.jumlah = jumlah_fisik
                    inventory.save()

                    # 2. Catat ke Stock Movement sebagai Audit Trail
                    StockMovement.objects.create(
                        barang=barang,
                        tipe='ADJUSTMENT',
                        jumlah=selisih, # Nilai bisa minus (hilang) atau plus (berlebih)
                        lokasi_asal=lokasi if selisih < 0 else None,
                        lokasi_tujuan=lokasi if selisih > 0 else None,
                        user=request.user,
                        keterangan=f"[ADJUSTMENT] Sistem: {stok_sistem}, Fisik: {jumlah_fisik}. Alasan: {keterangan}"
                    )
                    messages.success(request, f"Penyesuaian stok {barang.nama} berhasil dicatat!")
                else:
                    messages.info(request, "Jumlah fisik sama dengan sistem, tidak ada perubahan yang dicatat.")
                
                return redirect('stok_barang')
                
        except Exception as e:
            messages.error(request, f"Terjadi kesalahan saat memproses data: {e}")
            return redirect('adjustment_stok')

    # Method GET: Kirim dropdown Barang & Lokasi ke template
    barangs = Barang.objects.all().order_by('nama')
    lokasis = Lokasi.objects.all().order_by('nama_lokasi')
    
    return render(request, 'inventaris/operasional/adjustment.html', {
        'barangs': barangs,
        'lokasis': lokasis
    })
# =========================
# KONTROL PERGERAKAN STOK
# =========================

@login_required(login_url='/login/')
def kontrol_stok_view(request):


    movements = StockMovement.objects.select_related(
        'barang',
        'barang__kategori',
        'barang__supplier',
        'lokasi_asal',
        'lokasi_tujuan',
        'user'
    ).prefetch_related(
        'barang__inventory_set__lokasi'
    ).order_by(
        '-tanggal'
    )


    return render(
        request,
        'inventaris/transaksi/Kontrol_Pergerakan_Stok.html',
        {
            'movements': movements
        }
    )


# =========================
# VALIDASI OPERASIONAL
# =========================

@login_required(login_url='/login/')
def validasi_operasional_view(request):


    

    # CEK STOK MINUS

    stok_minus = Inventory.objects.filter(

        jumlah__lt=0

    )



    # CEK INVENTORY TANPA LOKASI

    inventory_tanpa_lokasi = Inventory.objects.filter(

        lokasi__isnull=True

    )



    # CEK MOVEMENT TANPA USER

    movement_tanpa_user = StockMovement.objects.filter(

        user__isnull=True

    )



    # CEK TRANSFER INVALID

    transfer_invalid = StockMovement.objects.filter(

    tipe="TRANSFER",

    lokasi_asal=F('lokasi_tujuan')

    )



    context = {


        'stok_minus': stok_minus,


        'inventory_tanpa_lokasi': inventory_tanpa_lokasi,


        'movement_tanpa_user': movement_tanpa_user,


        'transfer_invalid': transfer_invalid,


    }


    return render(

        request,

        'inventaris/transaksi/Validasi_Operasional.html'

    )
# ========================================================
# WMS MONITOR
# ========================================================
@login_required(login_url='/login/')
def wms_monitor_view(request):
    today = timezone.now().date()

    # --- 1. METRIK DASAR ---
    total_barang = Barang.objects.count()
    total_stok = Inventory.objects.aggregate(total=Sum('jumlah'))['total'] or 0
    total_lokasi = Lokasi.objects.count()
    stok_kritis = Inventory.objects.filter(jumlah__lte=F('barang__minimum_stock')).select_related('barang', 'lokasi').order_by('jumlah')

    # --- 2. METRIK HARIAN (Unit Barang) ---
    masuk_hari_ini = StockMovement.objects.filter(tipe='MASUK', tanggal__date=today).aggregate(tot=Sum('jumlah'))['tot'] or 0
    keluar_hari_ini = StockMovement.objects.filter(tipe='KELUAR', tanggal__date=today).aggregate(tot=Sum('jumlah'))['tot'] or 0
    transfer_hari_ini = StockMovement.objects.filter(tipe='TRANSFER', tanggal__date=today).aggregate(tot=Sum('jumlah'))['tot'] or 0

    # --- 3. AUDIT TRAIL / AKTIVITAS TERBARU ---
    aktivitas_terbaru = StockMovement.objects.select_related('barang', 'user').order_by('-tanggal')[:5]

    # --- 4. DATA GRAFIK (7 Hari Terakhir) ---
    tujuh_hari_lalu = today - timedelta(days=6)
    pergerakan_7_hari = StockMovement.objects.filter(tanggal__date__gte=tujuh_hari_lalu) \
        .annotate(tanggal_only=TruncDate('tanggal')) \
        .values('tanggal_only', 'tipe') \
        .annotate(total=Sum('jumlah'))

    # Siapkan wadah data kosong untuk 7 hari agar grafik tidak bolong
    data_masuk = { (tujuh_hari_lalu + timedelta(days=i)).strftime('%d %b'): 0 for i in range(7) }
    data_keluar = { (tujuh_hari_lalu + timedelta(days=i)).strftime('%d %b'): 0 for i in range(7) }

    # Isi data grafik dari database
    for item in pergerakan_7_hari:
        tgl_str = item['tanggal_only'].strftime('%d %b')
        if item['tipe'] == 'MASUK' and tgl_str in data_masuk:
            data_masuk[tgl_str] = item['total']
        elif item['tipe'] == 'KELUAR' and tgl_str in data_keluar:
            data_keluar[tgl_str] = item['total']

    # Convert ke list JSON untuk JavaScript
    grafik_label = json.dumps(list(data_masuk.keys()))
    grafik_masuk = json.dumps(list(data_masuk.values()))
    grafik_keluar = json.dumps(list(data_keluar.values()))

    return render(request, 'inventaris/wms/wms_monitor.html', {
        'total_barang': total_barang,
        'total_stok': total_stok,
        'total_lokasi': total_lokasi,
        'stok_kritis': stok_kritis,
        
        'masuk_hari_ini': masuk_hari_ini,
        'keluar_hari_ini': keluar_hari_ini,
        'transfer_hari_ini': transfer_hari_ini,
        
        'aktivitas_terbaru': aktivitas_terbaru,
        
        'grafik_label': grafik_label,
        'grafik_masuk': grafik_masuk,
        'grafik_keluar': grafik_keluar,
    })

# ========================================================
# STOK BARANG
# ========================================================
@login_required(login_url='/login/')
def stok_barang_view(request):
    query = request.GET.get('q', '')
    
    inventory_list = Inventory.objects.select_related(
        'barang', 'barang__kategori', 'lokasi'
    ).order_by('barang__nama')

    # Logika Pencarian
    if query:
        inventory_list = inventory_list.filter(
            Q(barang__nama__icontains=query) |
            Q(barang__sku__icontains=query) |
            Q(barang__kategori__nama__icontains=query) |
            Q(lokasi__nama_lokasi__icontains=query)
        )

    # Ambil list semua barang untuk fitur Autocomplete/Suggestions
    semua_barang = Barang.objects.all()

    # Logika Pagination
    paginator = Paginator(inventory_list, 10)
    page_number = request.GET.get('page')
    inventory = paginator.get_page(page_number)

    return render(request, 'inventaris/wms/stok_barang.html', {
        'inventory': inventory,
        'query': query,
        'semua_barang': semua_barang
    })

# ========================================================
# LOKASI STOK
# ========================================================
@login_required(login_url='/login/')
def lokasi_stok_view(request):
    query = request.GET.get('q', '')
    
    lokasi_list = Lokasi.objects.annotate(
        total_barang=Count('inventory__barang', distinct=True),
        total_unit=Sum('inventory__jumlah')
    ).order_by('nama_lokasi')

    # Logika Pencarian
    if query:
        lokasi_list = lokasi_list.filter(
            Q(nama_lokasi__icontains=query) |
            Q(kode_lokasi__icontains=query)
        )
        
    # Ambil list semua lokasi untuk fitur Autocomplete/Suggestions
    semua_lokasi = Lokasi.objects.all()

    # Logika Pagination
    paginator = Paginator(lokasi_list, 10)
    page_number = request.GET.get('page')
    lokasi_stok = paginator.get_page(page_number)

    return render(request, 'inventaris/wms/lokasi_stok.html', {
        'lokasi_stok': lokasi_stok,
        'query': query,
        'semua_lokasi': semua_lokasi
    })
# ========================================================
# LAPORAN & EXPORT
# ========================================================

@login_required(login_url='/login/')
def laporan_view(request):
    return render(request, 'inventaris/laporan/laporan_dashboard.html')


@login_required(login_url='/login/')
def export_inventory_csv(request):
    # Setup response agar diunduh sebagai file CSV
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="Laporan_Stok_Inventory.csv"'
    
    writer = csv.writer(response)
    # Header Tabel
    writer.writerow(['SKU', 'Nama Barang', 'Kategori', 'Lokasi', 'Jumlah', 'Satuan', 'Status Stok'])
    
    inventory = Inventory.objects.select_related('barang', 'barang__kategori', 'lokasi').all()
    
    for item in inventory:
        # Tentukan status stok
        if item.jumlah == 0:
            status = "Kosong"
        elif item.jumlah <= item.barang.minimum_stock:
            status = "Menipis"
        else:
            status = "Aman"
            
        writer.writerow([
            item.barang.sku,
            item.barang.nama,
            item.barang.kategori.nama if item.barang.kategori else '-',
            item.lokasi.nama_lokasi if item.lokasi else '-',
            item.jumlah,
            item.barang.satuan,
            status
        ])
        
    return response


@login_required(login_url='/login/')
def export_transaksi_csv(request):
    # Setup response agar diunduh sebagai file CSV
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="Laporan_Riwayat_Transaksi.csv"'
    
    writer = csv.writer(response)
    # Header Tabel
    writer.writerow(['Tanggal', 'Tipe Transaksi', 'Barang', 'Jumlah', 'Lokasi Asal', 'Lokasi Tujuan', 'User', 'Keterangan'])
    
    transaksi = StockMovement.objects.select_related('barang', 'lokasi_asal', 'lokasi_tujuan', 'user').order_by('-tanggal')
    
    for t in transaksi:
        # Format tanggal ke string yang rapi
        tanggal_str = timezone.localtime(t.tanggal).strftime('%Y-%m-%d %H:%M:%S')
        
        writer.writerow([
            tanggal_str,
            t.tipe,
            t.barang.nama,
            t.jumlah,
            t.lokasi_asal.nama_lokasi if t.lokasi_asal else '-',
            t.lokasi_tujuan.nama_lokasi if t.lokasi_tujuan else '-',
            t.user.username if t.user else 'Sistem',
            t.keterangan
        ])
        
    return response

# =========================
# PROFIL
# =========================

@login_required(login_url='/login/')
def profil_view(request):


    profile, created = UserProfile.objects.get_or_create(
        user=request.user
    )


    if request.method == "POST":


        form = UserProfileForm(

            request.POST,

            request.FILES,

            instance=profile

        )


        if form.is_valid():


            form.save()


            messages.success(

                request,

                "Profil berhasil diperbarui"

            )


            return redirect('dasbor')



    else:


        form = UserProfileForm(

            instance=profile

        )



    return render(

        request,

        'inventaris/user/profil.html',

        {

            'form': form,

            'profile': profile

        }

    )

# =========================
# USER MANAGEMENT
# =========================


@login_required(login_url='/login/')
def user_list_view(request):

    daftar_user = User.objects.all()


    return render(

        request,

        'inventaris/user/user_list.html',

        {
            'daftar_user': daftar_user
        }

    )





@login_required(login_url='/login/')
def tambah_user_view(request):


    if request.method == "POST":


        form = UserCreateForm(request.POST)


        if form.is_valid():


            form.save()


            messages.success(
                request,
                "User berhasil ditambahkan"
            )


            return redirect('user_list')



    else:


        form = UserCreateForm()



    return render(
    request,
    'inventaris/user/tambah_user.html',
    {
        'form': form
    }
)


@login_required(login_url='/login/')
def edit_user_view(request, id):


    user = get_object_or_404(
        User,
        id=id
    )



    if request.method == "POST":


        form = UserEditForm(

            request.POST,

            instance=user

        )


        if form.is_valid():


            form.save()



            messages.success(

                request,

                "Data user berhasil diperbarui"

            )



            return redirect('user_list')



    else:


        form = UserEditForm(

            instance=user

        )



    return render(

        request,

        'inventaris/user/edit_user.html',

        {

            'form':form

        }

    )

@login_required(login_url='/login/')
def hapus_user_view(request, id):

    user = get_object_or_404(
        User,
        id=id
    )


    if request.method == "POST":

        user.delete()


        messages.success(
            request,
            "User berhasil dihapus"
        )


    return redirect('user_list')


# =========================
# LOGIN
# =========================

def login_view(request):

    if request.method == "POST":

        username = request.POST.get('username')

        password = request.POST.get('password')


        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None:
            login(request, user)

            next_url = request.GET.get('next')

            if next_url:
                return redirect(next_url)

            return redirect('dasbor')

        else:
            messages.error(
                request,
                "Username atau password salah"
            )


    return render(
        request,
        'inventaris/login.html'
    )

# =========================
# LOGOUT
# =========================

def logout_akun(request):

    logout(request)

    return redirect('login')